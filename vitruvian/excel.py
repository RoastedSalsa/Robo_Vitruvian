"""The workbook: a filterable summary, one joint sheet per payload, and the inputs."""
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .drives import Drive
from .motors import Motor
from .selection import Selection
from .tables import JOINT_HEADER, joint_rows, summary_header, summary_rows

BOLD = Font(bold=True)
RED = Font(color="C00000")
FAIL = PatternFill("solid", fgColor="F8D7DA")
BAND = PatternFill("solid", fgColor="EEF1F5")
HEADER = PatternFill("solid", fgColor="D9DEE6")
TOP = Border(top=Side(style="medium"))


def sheet(wb: Workbook, title: str, header: list[str], rows: list[list], freeze: str = "A2",
          formats: dict[str, str] | None = None):
    """A sheet with a styled header row, frozen panes, filters and fitted columns."""
    ws = wb.create_sheet(title)
    ws.append(header)
    for cell in ws[1]:
        cell.font, cell.fill = BOLD, HEADER
    for row in rows:
        ws.append(row)
    ws.freeze_panes = freeze
    ws.auto_filter.ref = ws.dimensions
    for i, name in enumerate(header, 1):
        width = max(len(str(name)), *(len(f"{r[i - 1]}") for r in rows)) if rows else len(name)
        ws.column_dimensions[get_column_letter(i)].width = min(width + 2, 60)
        for cell in ws[get_column_letter(i)][1:]:
            if formats and name in formats:
                cell.number_format = formats[name]
    return ws


def write_workbook(path: Path, results: dict[float, list[Selection]], heights: list[float],
                   limits: list[float], catalog: list[Motor], options: list[list[Drive]],
                   cfg: dict) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    wb.remove(wb.worksheets[0])
    pct = "0%"

    # Summary: one row per payload and height; a heavy rule between payloads.
    rows = summary_rows(results, limits, cfg)
    ws = sheet(wb, "summary", summary_header(cfg), rows, freeze="C2",
               formats={"height m": "0.00", "drives kg": "0.00", "structure kg": "0.00",
                        "arm kg": "0.00", "max payload kg": "0.00"})
    for r, row in enumerate(rows, 2):
        if row[2] != "OK":
            for cell in ws[r]:
                cell.fill = FAIL
        if r > 2 and row[0] != rows[r - 3][0]:
            for cell in ws[r]:
                cell.border = TOP

    sheet(wb, "max payload", ["height m", "max payload kg"],
          [[h, round(kg, 2)] for h, kg in zip(heights, limits)], formats={"height m": "0.00"})

    # One sheet per payload: every joint at every height, alternate heights banded.
    fmt = {"x m": "0.000", "out kg": "0.00", "inertia kgm²": "0.0000", "alpha rad/s²": "0.0",
           "F hold N": "0.0", "F peak N": "0.0", "τ hold Nm": "0.00", "τ accel Nm": "0.00",
           "τ peak Nm": "0.00", "rated used": pct, "peak used": pct, "speed rad/s": "0.00",
           "speed req rad/s": "0.00", "Ø out mm": "0", "Ø max mm": "0", "height m": "0.00"}
    for kg, sels in results.items():
        rows = joint_rows(sels, cfg)
        ws = sheet(wb, f"joints {kg:g} kg", JOINT_HEADER, rows, freeze="C2", formats=fmt)
        n = len(cfg["joint"])
        for r, row in enumerate(rows, 2):
            if (r - 2) // n % 2:
                for cell in ws[r]:
                    cell.fill = BAND
            if row[-2] != "OK":
                for cell in ws[r]:
                    cell.font = RED

    m = cfg["motors"]
    sheet(wb, "motors", ["motor", "rated Nm", "peak Nm", "mass kg", "gear ratio",
                         "rotor g·cm²", "rated rpm", "no-load rpm", "rated V", "supply V",
                         "no-load rad/s at supply", "excluded", "estimated", "local", "source"],
          [[mo.name, mo.rated_Nm, mo.peak_Nm, mo.mass_kg, mo.ratio, mo.rotor_inertia_gcm2,
            mo.rated_rpm, mo.no_load_rpm, mo.rated_V, mo.supply_V, round(mo.speed_rad_s(0), 2),
            "yes" if mo.name in m["exclude"] else "", ", ".join(mo.estimated),
            "yes" if mo.local else "", mo.source]
           for mo in catalog])
    allowed_at: dict[str, list[str]] = {}
    drives: dict[str, Drive] = {}
    for spec, joint_options in zip(cfg["joint"], options):
        for d in joint_options:
            drives[d.name] = d
            allowed_at.setdefault(d.name, []).append(spec["name"])
    sheet(wb, "drive options", ["drive", "motor", "type", "extra ratio", "total ratio", "stages",
                                "efficiency", "rated Nm at joint", "peak Nm at joint", "mass kg",
                                "rotor at joint kgm²", "min Ø out mm", "allowed at"],
          [[d.name, d.motor.name, d.kind, d.ratio, d.total_ratio, d.stages,
            round(d.efficiency, 3), round(d.rated_Nm, 2), round(d.peak_Nm, 2), round(d.mass_kg, 3),
            round(d.reflected_inertia, 4), round(d.min_diameter_mm), " ".join(allowed_at[d.name])]
           for d in drives.values()])
    sheet(wb, "inputs", ["section", "key", "value"],
          [[section, key, str(value)] for section, table in cfg.items() if isinstance(table, dict)
           for key, value in table.items()] +
          [["drive", d["name"], str({k: v for k, v in d.items() if k != "name"})] for d in cfg["drive"]] +
          [["joint", j["name"], str({k: v for k, v in j.items() if k != "name"})] for j in cfg["joint"]])
    wb.save(path)
